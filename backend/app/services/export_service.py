import io
from typing import Dict, Any
from jinja2 import Template
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from app.schemas.resume import ResumeRead, ResumeBase


HTML_RESUME_TEMPLATE = """
<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="UTF-8">
  <title>{{ resume.personal_info.fullName }} - Resume</title>
  <style>
    @page { margin: 0.5in; size: letter; }
    body {
      font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Helvetica, Arial, sans-serif;
      color: #111827;
      line-height: 1.4;
      font-size: 10.5pt;
      margin: 0;
      padding: 20px;
    }
    h1 { font-size: 22pt; margin: 0 0 4px 0; text-transform: uppercase; letter-spacing: 0.5px; }
    .headline { font-size: 11pt; color: #4b5563; font-weight: 500; margin-bottom: 6px; }
    .contact-bar {
      display: flex;
      flex-wrap: wrap;
      gap: 12px;
      font-size: 9.5pt;
      color: #374151;
      border-bottom: 1.5px solid #111827;
      padding-bottom: 8px;
      margin-bottom: 14px;
    }
    .section-title {
      font-size: 12pt;
      font-weight: 700;
      text-transform: uppercase;
      letter-spacing: 1px;
      border-bottom: 1px solid #d1d5db;
      padding-bottom: 3px;
      margin-top: 14px;
      margin-bottom: 8px;
      color: #1f2937;
    }
    .entry { margin-bottom: 10px; }
    .entry-header { display: flex; justify-content: space-between; font-weight: 600; font-size: 10.5pt; }
    .entry-sub { display: flex; justify-content: space-between; font-style: italic; color: #4b5563; font-size: 9.5pt; margin-bottom: 3px; }
    ul { margin: 2px 0 6px 18px; padding: 0; }
    li { margin-bottom: 3px; }
    .skills-grid { font-size: 9.5pt; }
    .skill-category { margin-bottom: 4px; }
    .skill-cat-title { font-weight: 600; }
  </style>
</head>
<body>
  <h1>{{ resume.personal_info.fullName }}</h1>
  {% if resume.personal_info.headline %}
    <div class="headline">{{ resume.personal_info.headline }}</div>
  {% endif %}
  <div class="contact-bar">
    {% if resume.personal_info.email %}<span>📧 {{ resume.personal_info.email }}</span>{% endif %}
    {% if resume.personal_info.phone %}<span>📱 {{ resume.personal_info.phone }}</span>{% endif %}
    {% if resume.personal_info.location %}<span>📍 {{ resume.personal_info.location }}</span>{% endif %}
    {% if resume.personal_info.linkedin %}<span>🔗 {{ resume.personal_info.linkedin }}</span>{% endif %}
    {% if resume.personal_info.github %}<span>💻 {{ resume.personal_info.github }}</span>{% endif %}
  </div>

  {% if resume.personal_info.summary %}
    <div class="section-title">Professional Summary</div>
    <p style="margin: 0 0 10px 0;">{{ resume.personal_info.summary }}</p>
  {% endif %}

  {% if resume.experiences %}
    <div class="section-title">Work Experience</div>
    {% for exp in resume.experiences %}
      <div class="entry">
        <div class="entry-header">
          <span>{{ exp.title }}</span>
          <span>{{ exp.startDate }} — {{ 'Present' if exp.current else exp.endDate }}</span>
        </div>
        <div class="entry-sub">
          <span>{{ exp.company }}</span>
          <span>{{ exp.location }}</span>
        </div>
        {% if exp.bullets %}
          <ul>
            {% for bullet in exp.bullets %}
              <li>{{ bullet }}</li>
            {% endfor %}
          </ul>
        {% endif %}
      </div>
    {% endfor %}
  {% endif %}

  {% if resume.education %}
    <div class="section-title">Education</div>
    {% for edu in resume.education %}
      <div class="entry">
        <div class="entry-header">
          <span>{{ edu.institution }}</span>
          <span>{{ edu.startDate }} — {{ edu.endDate }}</span>
        </div>
        <div class="entry-sub">
          <span>{{ edu.degree }}{% if edu.fieldOfStudy %}, {{ edu.fieldOfStudy }}{% endif %}</span>
          {% if edu.gpa %}<span>GPA: {{ edu.gpa }}</span>{% endif %}
        </div>
      </div>
    {% endfor %}
  {% endif %}

  {% if resume.projects %}
    <div class="section-title">Key Projects</div>
    {% for proj in resume.projects %}
      <div class="entry">
        <div class="entry-header">
          <span>{{ proj.title }}</span>
          {% if proj.technologies %}<span>{{ proj.technologies | join(', ') }}</span>{% endif %}
        </div>
        <p style="margin: 2px 0 4px 0; font-size: 9.5pt;">{{ proj.description }}</p>
        {% if proj.bullets %}
          <ul>
            {% for bullet in proj.bullets %}
              <li>{{ bullet }}</li>
            {% endfor %}
          </ul>
        {% endif %}
      </div>
    {% endfor %}
  {% endif %}

  {% if resume.skills %}
    <div class="section-title">Technical Skills</div>
    <div class="skills-grid">
      {% for cat in resume.skills %}
        <div class="skill-category">
          <span class="skill-cat-title">{{ cat.category }}:</span> {{ cat.items | join(', ') }}
        </div>
      {% endfor %}
    </div>
  {% endif %}
</body>
</html>
"""


class ExportService:
    @staticmethod
    def render_html(resume: ResumeBase) -> str:
        template = Template(HTML_RESUME_TEMPLATE)
        return template.render(resume=resume)

    @staticmethod
    def generate_docx(resume: ResumeBase) -> bytes:
        """Generates a structured Microsoft Word (.docx) document."""
        doc = Document()

        # Document title / Name
        title_p = doc.add_paragraph()
        run = title_p.add_run(resume.personal_info.fullName)
        run.bold = True
        run.font.size = Pt(20)

        # Headline
        if resume.personal_info.headline:
            p = doc.add_paragraph()
            r = p.add_run(resume.personal_info.headline)
            r.font.size = Pt(11)
            r.font.color.rgb = RGBColor(100, 100, 100)

        # Contact info
        contacts = []
        if resume.personal_info.email: contacts.append(resume.personal_info.email)
        if resume.personal_info.phone: contacts.append(resume.personal_info.phone)
        if resume.personal_info.location: contacts.append(resume.personal_info.location)
        if contacts:
            doc.add_paragraph(" | ".join(contacts))

        # Summary
        if resume.personal_info.summary:
            h = doc.add_heading("Professional Summary", level=2)
            doc.add_paragraph(resume.personal_info.summary)

        # Experiences
        if resume.experiences:
            doc.add_heading("Work Experience", level=2)
            for exp in resume.experiences:
                p = doc.add_paragraph()
                r1 = p.add_run(f"{exp.title} - {exp.company}")
                r1.bold = True
                dates = f" ({exp.startDate} - {'Present' if exp.current else exp.endDate})"
                p.add_run(dates)
                for b in exp.bullets:
                    doc.add_paragraph(b, style="List Bullet")

        # Education
        if resume.education:
            doc.add_heading("Education", level=2)
            for edu in resume.education:
                p = doc.add_paragraph()
                p.add_run(f"{edu.institution} - {edu.degree}")

        # Skills
        if resume.skills:
            doc.add_heading("Technical Skills", level=2)
            for cat in resume.skills:
                p = doc.add_paragraph()
                r = p.add_run(f"{cat.category}: ")
                r.bold = True
                p.add_run(", ".join(cat.items))

        file_stream = io.BytesIO()
        doc.save(file_stream)
        return file_stream.getvalue()


export_service = ExportService()
